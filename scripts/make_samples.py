"""Generate sample .docx and .pdf syllabus documents used to verify the ingest
pipeline end-to-end. Run once:  python scripts/make_samples.py

Requires: pip install python-docx reportlab
"""
from __future__ import annotations

from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def make_docx() -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Calculus Notes — Differentiation (Extra Material)", level=0)
    doc.add_heading("Unit 6 — Introduction to Calculus", level=1)
    doc.add_paragraph(
        "These notes expand Unit 6 of the Mathematics syllabus with extra practice on derivatives. "
        "The derivative of a function measures how fast the function changes at a given point. "
        "If a car's distance s(t) changes with time t, then ds/dt is its instantaneous velocity."
    )
    doc.add_heading("The Power Rule", level=2)
    doc.add_paragraph(
        "For f(x) = x^n, the derivative is f'(x) = n x^(n-1). For example, the derivative of x^4 is 4x^3. "
        "Constant multiples stay as they are: the derivative of 5x^3 is 15x^2. Adding or subtracting terms "
        "works term by term, so the derivative of 2x^2 + 3x - 1 is 4x + 3."
    )
    doc.add_heading("Tangents and Rates of Change", level=2)
    doc.add_paragraph(
        "The derivative at x = a equals the gradient of the tangent line to the curve at that point. "
        "If f'(a) > 0 the function is rising there; if f'(a) < 0 it is falling. "
        "Practice: find the gradient of y = x^2 - 4x + 1 when x = 3. "
        "Since dy/dx = 2x - 4, at x = 3 the gradient is 2(3) - 4 = 2."
    )
    doc.add_heading("Exam-Style Question", level=2)
    doc.add_paragraph(
        "A ball is thrown upward so its height after t seconds is h(t) = 20t - 5t^2. "
        "Find the instantaneous velocity at t = 1 s. Differentiate to get h'(t) = 20 - 10t, "
        "then substitute t = 1: h'(1) = 10 m/s. The positive sign means the ball is still rising."
    )
    doc.save(ROOT / "data" / "math" / "calculus-notes.docx")


def make_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    out = ROOT / "data" / "physics" / "mechanics-extras.pdf"
    doc = SimpleDocTemplate(str(out), pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=16, spaceAfter=12)
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=12, spaceAfter=8)
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=10, leading=14, spaceAfter=6)

    story = [
        Paragraph("Mechanics Extras — Friction (Extra Material)", title),
        Paragraph("Unit 3 — Laws of Motion", h1),
        Paragraph(
            "These notes add worked examples on friction to Unit 3 of the Physics syllabus. "
            "Friction is the resistive force that acts parallel to a surface, opposing the motion "
            "or attempted motion of an object. Kinetic friction acts while an object slides.",
            body,
        ),
        Paragraph("Worked Example: Box on a Floor", h1),
        Paragraph(
            "A 20 kg box is pushed across a level floor with a horizontal force of 60 N. "
            "The friction between box and floor is 40 N. The net force is 60 - 40 = 20 N, "
            "so by Newton's second law the acceleration is a = F_net / m = 20 / 20 = 1 m/s^2.",
            body,
        ),
        Paragraph("When Does Friction Reverse? ", h1),
        Paragraph(
            "If the pushed force is removed, friction becomes the only horizontal force and brings "
            "the box to rest. Doubling the mass doubles the normal reaction and therefore roughly "
            "doubles the friction (friction is proportional to the normal force for sliding objects).",
            body,
        ),
        Paragraph("Exam-Style Question", h1),
        Paragraph(
            "A 5 kg crate rests on a ramp. Its weight is W = mg = 5 x 9.8 = 49 N. "
            "State which force balances the component of weight acting down the ramp, and explain "
            "what happens to that force as the ramp is tilted more steeply.",
            body,
        ),
    ]
    doc.build(story)


if __name__ == "__main__":
    (ROOT / "data/math").mkdir(parents=True, exist_ok=True)
    (ROOT / "data/physics").mkdir(parents=True, exist_ok=True)
    make_docx()
    make_pdf()
    print("Created data/math/calculus-notes.docx and data/physics/mechanics-extras.pdf")